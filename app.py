import os
import uuid
import logging
import html
import io
import re

import streamlit as st
import streamlit.components.v1 as components

# ============================================================
# OPTIONAL DEPENDENCIES WITH FALLBACK HANDLING
# ============================================================
try:
    from openai import OpenAI
except ImportError:
    OpenAI = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    from fpdf import FPDF
except ImportError:
    FPDF = None

# ============================================================
# PAGE CONFIG
# ============================================================
st.set_page_config(
    page_title="BAITHAK WITH AI",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded",
)

DEFAULT_MODEL = "gpt-6-luna"

# 4-Tier System Prompts
SYSTEM_PROMPTS = {
    "Naib Qasid AI": """
You are Baithak Naib Qasid AI, the intake dispatcher and workflow initiator.
Your core responsibilities:
1. Formally accept the incoming request from the user.
2. Acknowledge the task brief, catalog attachments, and format an official Docket/Routing Slip.
3. Pass the docket down to the Clerk with clear instructions on what needs to be drafted.
""",
    "Clerk AI": """
You are Baithak Clerk AI, the drafting and clerical documentation specialist.
Your core responsibilities:
1. Review the docket sent by Naib Qasid.
2. Draft clear, formal, executive-ready emails, memos, notices, or markdown tables.
3. Prepare content cleanly so it can be exported to PDF, DOCX, TXT, or Excel.
4. Pass the prepared draft to the Superintendent for verification.
""",
    "Superintendent AI": """
You are Baithak Superintendent AI, the quality auditor and administrative verifier.
Your core responsibilities:
1. Audit the draft submitted by the Clerk for accuracy, tone, completeness, and structure.
2. Provide a clear verification audit note highlighting corrections or confirming readiness.
3. Forward the verified file with recommendations to the Section Officer for final executive sign-off.
""",
    "Section Officer AI": """
You are Baithak Section Officer AI, the final approving authority.
Your core responsibilities:
1. Review the entire file (Docket -> Clerk Draft -> Superintendent Audit).
2. Issue the final sanction/verdict: **SANCTIONED & APPROVED**, **APPROVED WITH CONDITIONS**, or **REJECTED**.
3. Authorize the immediate release and generation of downloadable files (PDF, Word, Excel, TXT).
"""
}

logging.basicConfig(level=logging.ERROR)
logger = logging.getLogger("baithak_ai")


# ============================================================
# SECRETS & CLIENT INIT
# ============================================================
def get_secret(name, default=None):
    try:
        value = st.secrets.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    except Exception:
        pass
    value = os.getenv(name)
    return str(value).strip() if value and str(value).strip() else default


API_KEY = get_secret("OPENAI_API_KEY")
MODEL = get_secret("OPENAI_MODEL", DEFAULT_MODEL)


@st.cache_resource(show_spinner=False)
def get_client(api_key):
    if not api_key or OpenAI is None:
        return None
    try:
        return OpenAI(api_key=api_key, timeout=60.0, max_retries=1)
    except Exception:
        logger.exception("OpenAI client initialization failed.")
        return None


client = get_client(API_KEY)


# ============================================================
# SESSION STATE INITIALIZATION
# ============================================================
if "conversations" not in st.session_state:
    first_id = str(uuid.uuid4())
    st.session_state.conversations = {
        first_id: {"title": "New chat", "messages": [], "documents": []}
    }
    st.session_state.active_conversation = first_id

if "active_conversation" not in st.session_state:
    st.session_state.active_conversation = next(iter(st.session_state.conversations))

if st.session_state.active_conversation not in st.session_state.conversations:
    st.session_state.active_conversation = next(iter(st.session_state.conversations))

if "notice" not in st.session_state:
    st.session_state.notice = ""
if "mode" not in st.session_state:
    st.session_state.mode = "Demo Mode"
if "ai_role" not in st.session_state:
    st.session_state.ai_role = "Full Multi-Agent Hierarchy (Naib Qasid -> Clerk -> Superintendent -> Section Officer)"


def new_chat():
    chat_id = str(uuid.uuid4())
    st.session_state.conversations[chat_id] = {
        "title": "New chat", "messages": [], "documents": []
    }
    st.session_state.active_conversation = chat_id
    st.session_state.notice = ""


def clear_current_chat():
    chat_id = st.session_state.active_conversation
    st.session_state.conversations[chat_id] = {
        "title": "New chat", "messages": [], "documents": []
    }
    st.session_state.notice = ""


def delete_all_chats():
    chat_id = str(uuid.uuid4())
    st.session_state.conversations = {
        chat_id: {"title": "New chat", "messages": [], "documents": []}
    }
    st.session_state.active_conversation = chat_id
    st.session_state.notice = ""


def update_title(conversation, prompt):
    if conversation["title"] == "New chat":
        title = " ".join(prompt.split())
        conversation["title"] = title[:35].rstrip() + ("..." if len(title) > 35 else "")


# ============================================================
# EXPORT HELPERS (PDF, DOCX, TXT, EXCEL)
# ============================================================
def generate_txt(text: str) -> bytes:
    return text.encode("utf-8")


def generate_pdf(text: str) -> bytes:
    if FPDF is None:
        raise RuntimeError("PDF generator library `fpdf2` is not installed.")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", size=11)
    clean_text = re.sub(r'[*_`#]', '', text)
    for line in clean_text.split("\n"):
        pdf.multi_cell(0, 7, txt=line)
    buffer = io.BytesIO()
    pdf.output(buffer)
    return buffer.getvalue()


def generate_docx(text: str) -> bytes:
    if Document is None:
        raise RuntimeError("Word generator library `python-docx` is not installed.")
    doc = Document()
    doc.add_heading("Baithak Official File Export", level=1)
    lines = text.split("\n")
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        elif stripped:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def parse_markdown_tables(text: str):
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    tables = []
    current_table = []
    for line in lines:
        if line.startswith("|") and line.endswith("|"):
            current_table.append(line)
        else:
            if len(current_table) >= 3:
                tables.append(current_table)
            current_table = []
    if len(current_table) >= 3:
        tables.append(current_table)

    dataframes = []
    for table_lines in tables:
        rows = []
        for row_str in table_lines:
            if re.match(r'^\|[\s\:\-|-]+\|$', row_str):
                continue
            cols = [col.strip() for col in row_str.split("|")[1:-1]]
            rows.append(cols)
        if len(rows) >= 2 and pd is not None:
            df = pd.DataFrame(rows[1:], columns=rows[0])
            dataframes.append(df)
    return dataframes


def generate_excel_from_tables(text: str) -> bytes:
    if pd is None:
        raise RuntimeError("Pandas or openpyxl missing.")
    dfs = parse_markdown_tables(text)
    if not dfs:
        raise ValueError("No table found to export.")
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        for idx, df in enumerate(dfs, start=1):
            df.to_excel(writer, sheet_name=f"Sheet_{idx}", index=False)
    return buffer.getvalue()


# ============================================================
# LLM ENGINE & 4-TIER HIERARCHY WORKFLOW
# ============================================================
def call_llm(messages, system_instruction):
    if client is None:
        return None
    try:
        response = client.responses.create(
            model=MODEL,
            instructions=system_instruction,
            input=messages,
        )
        answer = getattr(response, "output_text", None)
        return answer.strip() if answer else None
    except Exception:
        logger.exception("API call failed.")
        return None


def demo_reply():
    return (
        "📨 **[1. Naib Qasid - Intake & Docket Creation]**\n"
        "• Request Received: Formal Request Dispatch\n"
        "• Status: Docket created and handed over to Clerk for drafting.\n\n"
        "---\n\n"
        "📝 **[2. Clerk AI - Draft Preparation]**\n\n"
        "**Subject:** Official Communication Brief\n\n"
        "Dear Authority,\n\n"
        "The file draft has been prepared as per guidelines:\n\n"
        "| Section | Detail | Status |\n"
        "| :--- | :--- | :--- |\n"
        "| File No. | BTH-2026-09 | Formulated |\n"
        "| Content | Requirements & Action Plan | Ready for Verification |\n\n"
        "---\n\n"
        "🔍 **[3. Superintendent AI - Audit & Verification]**\n"
        "• Audit Verdict: **VERIFIED**\n"
        "• Verification Note: Structure, table formatting, and clarity verified. Submitted to Section Officer.\n\n"
        "---\n\n"
        "🏛️ **[4. Section Officer AI - Executive Sanction & Sign-off]**\n"
        "• Final Verdict: **SANCTIONED & APPROVED FOR RELEASE**\n"
        "• Authorization: Granted. File is cleared for PDF, Word, Excel, and TXT export.\n\n"
        "_Demo Mode Active. Configure `OPENAI_API_KEY` for live AI generation._"
    )


def execute_hierarchy_workflow(messages):
    selected_role = st.session_state.ai_role

    if client is None:
        st.session_state.mode = "Demo Mode"
        st.session_state.notice = "Demo Mode · Add OPENAI_API_KEY in Streamlit Secrets."
        return demo_reply()

    try:
        if selected_role == "Full Multi-Agent Hierarchy (Naib Qasid -> Clerk -> Superintendent -> Section Officer)":
            # 1. Naib Qasid Accepts & Routes Brief
            nq_res = call_llm(messages
